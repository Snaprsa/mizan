"""Export routes — CSV, XLSX, and optional PDF."""
import os
import csv
import io
import uuid
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from ..core.database import get_db
from ..core.config import settings
from ..services.extraction_service import ExtractionService

router = APIRouter(tags=["export"])


@router.get("/cases/{case_id}/export/csv")
async def export_csv(case_id: str, db: AsyncSession = Depends(get_db)):
    svc = ExtractionService(db)
    comps = await svc.get_comps(case_id)
    if not comps:
        raise HTTPException(status_code=404, detail="No comparables found")

    output = io.StringIO()
    writer = csv.writer(output)
    headers = [
        "ID", "Source URL", "Source Type", "Property Title", "City", "District",
        "Subdistrict", "Latitude", "Longitude", "Google Maps Link",
        "Property Category", "Property Subtype",
        "Evidence Type", "Asking Price (SAR)", "Transacted Price (SAR)", "Currency",
        "Land Area (sqm)", "BUA (sqm)", "Price/sqm", "Listing Date",
        "Transaction Date", "Tenure", "Reliability Score", "Is Invalid", "Is Outlier", "Notes",
    ]
    writer.writerow(headers)

    for c in comps:
        land_area = ""
        bua = ""
        for a in (c.areas or []):
            if isinstance(a, dict):
                if a.get("area_type") in ("land_area", "plot_area"):
                    land_area = a.get("normalized_value", "")
                elif a.get("area_type") == "bua":
                    bua = a.get("normalized_value", "")
        gmaps = f"https://maps.google.com/?q={c.latitude},{c.longitude}" if c.latitude and c.longitude else ""
        writer.writerow([
            c.id, c.source_url or "", c.source_type, c.property_title or "",
            c.city or "", c.district or "", c.subdistrict or "",
            c.latitude or "", c.longitude or "", gmaps,
            c.property_category or "", c.property_subtype or "",
            c.evidence_type or "", c.asking_price or "", c.transacted_price or "",
            c.currency or "SAR", land_area, bua,
            c.price_per_sqm or "", c.listing_date or "", c.transaction_date or "",
            c.tenure or "", c.reliability_score or "", c.is_invalid, c.is_outlier,
            c.notes or "",
        ])

    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=comps_{case_id[:8]}.csv"},
    )


@router.get("/cases/{case_id}/export/xlsx")
async def export_xlsx(case_id: str, db: AsyncSession = Depends(get_db)):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    svc = ExtractionService(db)
    comps = await svc.get_comps(case_id)
    if not comps:
        raise HTTPException(status_code=404, detail="No comparables found")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparables"

    headers = [
        "ID", "Source URL", "Source Type", "Property Title", "City", "District",
        "Subdistrict", "Lat", "Lng", "Google Maps Link", "Category", "Subtype",
        "Evidence Type", "Asking Price (SAR)", "Transacted Price (SAR)", "Currency",
        "Land Area (sqm)", "BUA (sqm)", "Price/sqm", "Listing Date",
        "Transaction Date", "Tenure", "Reliability", "Invalid", "Outlier", "Notes",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, c in enumerate(comps, 2):
        land_area = ""
        bua = ""
        for a in (c.areas or []):
            if isinstance(a, dict):
                if a.get("area_type") in ("land_area", "plot_area"):
                    land_area = a.get("normalized_value", "")
                elif a.get("area_type") == "bua":
                    bua = a.get("normalized_value", "")
        gmaps = f"https://maps.google.com/?q={c.latitude},{c.longitude}" if c.latitude and c.longitude else ""
        row_data = [
            c.id[:8], c.source_url or "", c.source_type, c.property_title or "",
            c.city or "", c.district or "", c.subdistrict or "",
            c.latitude or "", c.longitude or "", gmaps,
            c.property_category or "", c.property_subtype or "",
            c.evidence_type or "", c.asking_price, c.transacted_price,
            c.currency or "SAR", land_area, bua,
            c.price_per_sqm, c.listing_date or "", c.transaction_date or "",
            c.tenure or "", c.reliability_score, c.is_invalid, c.is_outlier,
            c.notes or "",
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=comps_{case_id[:8]}.xlsx"},
    )
