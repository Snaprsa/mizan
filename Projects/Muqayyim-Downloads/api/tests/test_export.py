"""
Tests for export functionality — CSV, XLSX, PDF.
Tests the reporting modules directly (no DB needed).
"""

import io
import csv
import sys
import os
from datetime import datetime
from types import SimpleNamespace

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))


def _make_mock_case():
    """Create a mock Case object for export testing."""
    return SimpleNamespace(
        id="00000000-0000-0000-0000-000000000001",
        title="Test Case — Riyadh Villas",
        description="Sample case for export testing",
        status="in_progress",
        created_at=datetime(2026, 1, 15, 12, 0, 0),
    )


def _make_mock_comp(
    asking_price=3500000,
    area_sqm=450,
    city="Riyadh",
    district="Al Nakheel",
    property_type="villa",
):
    """Create a mock Comparable object."""
    price_per_sqm = round(asking_price / area_sqm, 2) if area_sqm else None
    return SimpleNamespace(
        id="00000000-0000-0000-0000-000000000010",
        case_id="00000000-0000-0000-0000-000000000001",
        source_id="00000000-0000-0000-0000-000000000020",
        status="active",
        reliability_score=3,
        property_title=f"Villa in {district}",
        city=city,
        district=district,
        property_type=property_type,
        evidence_type="listing_asking",
        asking_price=asking_price,
        transacted_price=None,
        currency="SAR",
        primary_area_sqm=area_sqm,
        price_per_sqm=price_per_sqm,
        listing_date="2025-01-15",
        transaction_date=None,
        tenure="freehold",
        latitude=24.7136,
        longitude=46.6753,
        notes=None,
        created_at=datetime(2026, 1, 15, 12, 0, 0),
    )


def _make_mock_source(url="https://aqar.fm/listing/12345"):
    """Create a mock Source object."""
    return SimpleNamespace(
        id="00000000-0000-0000-0000-000000000020",
        url=url,
        source_type="url",
    )


def _make_mock_fields():
    """Create mock ExtractedField objects."""
    return [
        SimpleNamespace(
            id="00000000-0000-0000-0000-000000000030",
            field_name="asking_price",
            raw_value="3,500,000 ريال",
            normalized_value="3500000",
            display_value="3,500,000",
            confidence=0.95,
            overrides=[],
        ),
        SimpleNamespace(
            id="00000000-0000-0000-0000-000000000031",
            field_name="primary_area_sqm",
            raw_value="450 متر مربع",
            normalized_value="450",
            display_value="450",
            confidence=0.9,
            overrides=[],
        ),
    ]


class TestCSVExport:
    """Tests for CSV export generation."""

    def test_generates_valid_csv(self):
        from reporting.csv_export import generate_csv

        case = _make_mock_case()
        comp = _make_mock_comp()
        source = _make_mock_source()
        fields = _make_mock_fields()

        result = generate_csv(case, [(comp, fields, source)])

        assert isinstance(result, io.BytesIO)
        content = result.getvalue().decode("utf-8-sig")
        assert len(content) > 0

        # Parse CSV and validate
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 2  # Header + 1 data row

        header = rows[0]
        assert "City" in header
        assert "Asking Price (SAR)" in header
        assert "Area (sqm)" in header

    def test_csv_contains_correct_data(self):
        from reporting.csv_export import generate_csv

        case = _make_mock_case()
        comp = _make_mock_comp(asking_price=2000000, area_sqm=300, city="Jeddah")
        source = _make_mock_source(url="https://example.com")
        fields = _make_mock_fields()

        result = generate_csv(case, [(comp, fields, source)])
        content = result.getvalue().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        header = rows[0]
        data_row = rows[1]
        city_idx = header.index("City")
        price_idx = header.index("Asking Price (SAR)")

        assert data_row[city_idx] == "Jeddah"
        assert data_row[price_idx] == "2000000"

    def test_csv_multiple_comps(self):
        from reporting.csv_export import generate_csv

        case = _make_mock_case()
        comp1 = _make_mock_comp(city="Riyadh")
        comp2 = _make_mock_comp(city="Jeddah", asking_price=4000000)
        source = _make_mock_source()
        fields = _make_mock_fields()

        result = generate_csv(case, [(comp1, fields, source), (comp2, fields, source)])
        content = result.getvalue().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        assert len(rows) == 3  # Header + 2 data rows

    def test_csv_empty_comps(self):
        from reporting.csv_export import generate_csv

        case = _make_mock_case()
        result = generate_csv(case, [])
        content = result.getvalue().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        assert len(rows) == 1  # Header only


class TestXLSXExport:
    """Tests for XLSX export generation."""

    def test_generates_valid_xlsx(self):
        from reporting.xlsx_export import generate_xlsx

        case = _make_mock_case()
        comp = _make_mock_comp()
        source = _make_mock_source()
        fields = _make_mock_fields()

        result = generate_xlsx(case, [(comp, fields, source)])

        assert isinstance(result, io.BytesIO)
        content = result.getvalue()
        assert len(content) > 0
        # XLSX files start with PK (zip signature)
        assert content[:2] == b"PK"

    def test_xlsx_readable_by_openpyxl(self):
        from reporting.xlsx_export import generate_xlsx
        import openpyxl

        case = _make_mock_case()
        comp = _make_mock_comp()
        source = _make_mock_source()
        fields = _make_mock_fields()

        result = generate_xlsx(case, [(comp, fields, source)])
        wb = openpyxl.load_workbook(result)
        ws = wb.active

        # Should have title row (1), blank (2), header row (3), data row (4+)
        assert ws.max_row >= 4
        # Headers are in row 3 (row 1 is a merged title row)
        headers = [cell.value for cell in ws[3] if cell.value]
        assert any("City" in str(h) for h in headers)


class TestPDFExport:
    """Tests for PDF summary export generation."""

    def test_generates_valid_pdf(self):
        from reporting.pdf_summary import generate_pdf_summary

        case = _make_mock_case()
        comp = _make_mock_comp()
        source = _make_mock_source()
        fields = _make_mock_fields()

        result = generate_pdf_summary(case, [(comp, fields, source)])

        assert isinstance(result, io.BytesIO)
        content = result.getvalue()
        assert len(content) > 0
        # PDF files start with %PDF
        assert content[:4] == b"%PDF"
