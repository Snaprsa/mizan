"""
XLSX export for comparables with formatting.
"""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# Confidence color thresholds
HIGH_CONF = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")    # Green
MED_CONF = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")     # Yellow
LOW_CONF = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")      # Red

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

COLUMNS = [
    ("Source", 30),
    ("Type", 10),
    ("Status", 10),
    ("Reliability", 10),
    ("Title", 30),
    ("City", 15),
    ("District", 15),
    ("Property Type", 15),
    ("Evidence", 15),
    ("Asking Price (SAR)", 18),
    ("Transacted Price (SAR)", 18),
    ("Area (sqm)", 12),
    ("Price/sqm", 12),
    ("Listing Date", 14),
    ("Transaction Date", 14),
    ("Tenure", 12),
    ("Lat", 12),
    ("Lng", 12),
    ("Notes", 40),
]


def generate_xlsx(case: Any, comps_with_fields: list[tuple]) -> io.BytesIO:
    """Generate formatted XLSX from case comparables."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparables"

    # Title row
    ws.merge_cells("A1:S1")
    ws["A1"] = f"Comparable Summary — {case.title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    # Header row
    for col_idx, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=3, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = width

    # Data rows
    for row_idx, (comp, fields, source) in enumerate(comps_with_fields, 4):
        data = [
            source.url or source.file_name if source else "",
            source.source_type if source else "",
            comp.status,
            comp.reliability_score,
            comp.property_title or "",
            comp.city or "",
            comp.district or "",
            comp.property_type or "",
            comp.evidence_type or "",
            comp.asking_price,
            comp.transacted_price,
            comp.primary_area_sqm,
            comp.price_per_sqm,
            comp.listing_date or "",
            comp.transaction_date or "",
            comp.tenure or "",
            comp.latitude,
            comp.longitude,
            comp.notes or "",
        ]

        for col_idx, val in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = THIN_BORDER

            # Format currency columns
            if col_idx in (10, 11, 13) and isinstance(val, (int, float)):
                cell.number_format = '#,##0'

        # Color-code row based on average field confidence
        if fields:
            avg_conf = sum(f.confidence for f in fields) / len(fields)
            fill = HIGH_CONF if avg_conf >= 0.8 else MED_CONF if avg_conf >= 0.5 else LOW_CONF
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill

    # Freeze header
    ws.freeze_panes = "A4"

    # Auto-filter
    ws.auto_filter.ref = f"A3:S{3 + len(comps_with_fields)}"

    # Write to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
